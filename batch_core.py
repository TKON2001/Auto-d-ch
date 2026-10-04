"""Ordered Excel download queue with atomic progress and portable project folders."""
from __future__ import annotations
import csv, hashlib, json, math, os, re, subprocess, sys, unicodedata
from pathlib import Path
from urllib.parse import urlsplit, parse_qs, unquote
from storage import atomic_json

class StopAfterVideo(Exception): pass

def normalize(value):
    text=unicodedata.normalize('NFD',str(value or '').strip().lower().replace('đ','d'))
    return ''.join(c for c in text if not unicodedata.combining(c))

def safe_name(value):
    value=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(value)).strip(' .')[:60]
    if not value:value='video'
    if value.upper().split('.')[0] in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}:value='_'+value
    return value

def extract_url(value):
    match=re.search(r'https?://[^\s<>"\u3000]+',str(value or ''))
    if not match:raise ValueError('Không có link http/https.')
    url=match.group(0).rstrip('.,;，。）)')
    host=urlsplit(url).hostname
    if not host:raise ValueError('Link không hợp lệ.')
    p=urlsplit(url)
    if host in ('douyin.com','www.douyin.com'):
        modal=parse_qs(p.query).get('modal_id',[''])[0]
        if modal.isdigit():url='https://www.douyin.com/video/'+modal
    return url

def local_source(value, excel):
    value=str(value).strip().strip('"')
    if value.lower().startswith('file:'):
        parsed=urlsplit(value);value=unquote(parsed.path)
        if parsed.netloc and parsed.netloc.lower()!='localhost':value='//'+parsed.netloc+value
        elif os.name=='nt' and re.match(r'^/[A-Za-z]:',value):value=value[1:]
    path=Path(os.path.expandvars(value)).expanduser()
    if not path.is_absolute():path=Path(excel).resolve().parent/path
    return str(path.resolve())

def copy_local_video(job, folder, control, log):
    import shutil
    source=Path(job['local_path'])
    if not source.is_file():raise FileNotFoundError('Không tìm thấy video trên máy: '+str(source))
    if source.suffix.lower() not in {'.mp4','.mkv','.mov','.webm','.avi','.m4v'}:raise ValueError('File không có định dạng video được hỗ trợ: '+source.name)
    target=Path(folder)/('source'+source.suffix.lower());temp=target.with_suffix(target.suffix+'.part')
    total=source.stat().st_size;done=0;last=-1
    try:
        with source.open('rb') as src,temp.open('wb') as dst:
            before=os.fstat(src.fileno())
            while True:
                control.check();chunk=src.read(1024*1024)
                if not chunk:break
                dst.write(chunk);done+=len(chunk);pct=int(done*100/total) if total else 100
                if pct//10!=last:log(f'Sao chép {source.name}: {pct}%');last=pct//10
            after=os.fstat(src.fileno())
            if (before.st_size,before.st_mtime_ns)!=(after.st_size,after.st_mtime_ns):raise RuntimeError('Video gốc thay đổi trong khi sao chép. Hãy thử lại.')
        shutil.copystat(source,temp);os.replace(temp,target)
    finally:temp.unlink(missing_ok=True)
    return target

def read_queue(path,sheet=None):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path,data_only=False);ws=wb[sheet] if sheet else wb.active
        name=ws.title
        values=[]
        for cells in ws.iter_rows():
            values.append([c.hyperlink.target if c.hyperlink and c.hyperlink.target else c.value for c in cells])
        wb.close()
    elif path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:values=list(csv.reader(f))
        name='CSV'
    else:raise ValueError('Chọn file .xlsx hoặc CSV UTF-8. File .xls cần lưu lại thành .xlsx.')
    if not values:raise ValueError('Bảng trống.')
    headers=[normalize(v) for v in values[0]]
    def column(aliases):return next((i for i,v in enumerate(headers) if v in aliases),None)
    link=column({'link','url','link video','video url'});order=column({'stt','thu tu','order'});title=column({'ten','ten video','title','name'})
    if link is None:raise ValueError('Dòng đầu cần cột Link hoặc URL.')
    jobs=[]
    for row,valuesrow in enumerate(values[1:],2):
        def val(col):return valuesrow[col] if col is not None and col<len(valuesrow) else None
        raw=val(link)
        if raw is None or not str(raw).strip():continue
        err='';url='';local=''
        if re.search(r'https?://',str(raw),re.I):
            try:url=extract_url(raw)
            except ValueError as exc:err=str(exc)
        else:local=local_source(raw,path)
        number=row-1
        if order is not None:
            try:
                number=float(val(order));assert math.isfinite(number) and number>0
            except (TypeError,ValueError,AssertionError):err=(err+'; ' if err else '')+'STT cần là số dương.';number=float('inf')
        item={'row':row,'order':number if math.isfinite(number) else None,'url':url,'local_path':local,'raw':str(raw),'name':str(val(title) or f'video_{row-1}'),'error':err,'status':'pending'}
        identity=json.dumps([name,row,item['order'],item['url'],item['raw'],item['name']],ensure_ascii=False)
        item['id']=hashlib.sha256(identity.encode()).hexdigest()[:12];jobs.append(item)
    if not jobs:raise ValueError('Không có dòng chứa link.')
    jobs.sort(key=lambda j:(j['order'] if j['order'] is not None else float('inf'),j['row']))
    for i,job in enumerate(jobs,1):job['folder']=f"{i:04d}_{safe_name(job['name'])}_{job['id'][:6]}"
    return name,jobs

def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for data in iter(lambda:f.read(1024*1024),b''):h.update(data)
    return h.hexdigest()

def download_video(job,folder,cookies,control,log):
    from components import find_binary
    ffmpeg=find_binary('ffmpeg')
    args=[sys.executable,'-m','yt_dlp','--no-playlist','--no-colors','--newline','--progress','--progress-delta','0.5','--socket-timeout','30','--retries','2','--fragment-retries','2','--windows-filenames','--print','after_move:filepath','--no-simulate','-f','bv*+ba/b','--merge-output-format','mp4','-o',str(Path(folder)/'source.%(ext)s')]
    if ffmpeg:args+=['--ffmpeg-location',str(Path(ffmpeg).parent)]
    if cookies:args+=['--cookies',str(cookies)]
    args+=['--',job['url']]
    proc=subprocess.Popen(args,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,encoding='utf-8',errors='replace',creationflags=0x08000000 if os.name=='nt' else 0)
    tail=[];final=None
    try:
        control.register(proc)
        for line in proc.stdout:
            text=line.rstrip();tail.append(text);tail=tail[-15:];log(text)
            candidate=Path(text)
            try:
                if candidate.is_file() and candidate.resolve().parent==Path(folder).resolve():final=candidate.resolve()
            except OSError:pass
        if proc.wait():raise RuntimeError('\n'.join(tail)[-2500:] or 'yt-dlp không tải được video.')
    finally:
        control.unregister(proc)
        if proc.poll() is None:proc.terminate();proc.wait()
    if not final:raise RuntimeError('Không tìm thấy video sau khi tải; kiểm tra nhật ký.')
    return final

class BatchQueue:
    def __init__(self,excel,destination,sheet=None):
        self.excel=Path(excel).resolve();self.sheet,self.jobs=read_queue(excel,sheet)
        identity=hashlib.sha256(json.dumps([self.excel.name,self.sheet,[j['id'] for j in self.jobs]],ensure_ascii=False).encode()).hexdigest()[:10]
        self.root=Path(destination).resolve()/f'Du_an_{safe_name(self.excel.stem)}_{identity}'
        self.root.mkdir(parents=True,exist_ok=True);self.manifest=self.root/'project.json';self.result=self.root/'Tien_do.xlsx';self.warning='';self.stop_requested=False
        if self.manifest.exists():
            old=json.loads(self.manifest.read_text(encoding='utf-8'));mapping={j['id']:j for j in old.get('videos',[])}
            for j in self.jobs:
                if j['id'] in mapping:j.update({k:v for k,v in mapping[j['id']].items() if k in {'status','video','project','sha256','info','error','translation_status','translation_output','translation_error','translation_signature','translation_sha256'}})
        self.persist()
    def persist(self):
        atomic_json(self.manifest,{'version':1,'type':'batch_video_project','excel':self.excel.name,'sheet':self.sheet,'videos':self.jobs})
    def save_excel(self):
        # Runtime feature of the user's Python app: preserve original workbook on disk.
        from openpyxl import load_workbook, Workbook
        from openpyxl.styles import PatternFill
        if self.excel.suffix.lower()=='.xlsx':wb=load_workbook(self.excel);ws=wb[self.sheet]
        else:
            wb=Workbook();ws=wb.active;ws.title='CSV'
            with self.excel.open(encoding='utf-8-sig',newline='') as f:
                for values in csv.reader(f):
                    ws.append(values)
                    for cell in ws[ws.max_row]:cell.data_type='s'
        last=ws.max_column
        for i,label in enumerate(['Frame: trạng thái','Frame: video','Frame: lỗi','Frame: dịch','Frame: video Việt','Frame: lỗi dịch'],last+1):ws.cell(1,i,label)
        for j in self.jobs:
            r=j['row'];status=j['status'];ws.cell(r,last+1,{'completed':('Đã nhập' if j.get('local_path') else 'Đã tải'),'failed':'Lỗi','downloading':'Đang tải','pending':'Chờ'}.get(status,status));ws.cell(r,last+2,j.get('video',''));ws.cell(r,last+3,j.get('error',''))
            for offset,key in [(4,'translation_status'),(5,'translation_output'),(6,'translation_error')]:ws.cell(r,last+offset,j.get(key,''))
            for col in range(last+1,last+7):ws.cell(r,col).data_type='s'
            if status in ('completed','failed'):
                color='FFC7CE' if status=='failed' or j.get('translation_status')=='failed' else ('FFEB9C' if j.get('translation_status')=='translating' else 'C6EFCE')
                fill=PatternFill('solid',fgColor=color)
                for col in range(1,last+7):ws.cell(r,col).fill=fill
        tmp=self.result.with_name('Tien_do.tmp.xlsx')
        try:wb.save(tmp);os.replace(tmp,self.result);self.warning=''
        except OSError as exc:self.warning='Không lưu được Excel kết quả. Đóng file đang mở rồi bấm Lưu Excel: '+str(exc)
        finally:wb.close();tmp.unlink(missing_ok=True)
        return self.warning
    def run(self,control,update=lambda *_:None,log=lambda _:None,downloader=download_video,probe=None):
        if probe is None:
            from video_translator import probe
        def publish(job):
            self.persist();warning=self.save_excel();update(job.copy())
            if warning:log(warning)
        for job in self.jobs:
            control.check()
            if self.stop_requested:break
            folder=self.root/job['folder'];folder.mkdir(exist_ok=True)
            if job['status']=='completed':
                existing=self.root/job.get('video','')
                local=Path(job['local_path']) if job.get('local_path') else None
                source_unchanged=local is None or not local.is_file() or digest(local)==job.get('sha256')
                if source_unchanged and existing.is_file() and digest(existing)==job.get('sha256'):
                    if not (folder/(folder.name+'.frameproj')).is_file():pack_video_project(folder)
                    publish(job);continue
                job.update(status='pending',error='')
            if (not job['url'] and not job.get('local_path')) or job['order'] is None:
                job.update(status='failed');publish(job);continue
            try:
                job.update(status='downloading',error='');publish(job)
                if job.get('local_path'):video=copy_local_video(job,folder,control,log).resolve()
                else:video=Path(downloader(job,folder,None if not getattr(self,'cookies',None) else self.cookies,control,log)).resolve()
                if not video.is_relative_to(folder.resolve()) or not video.is_file():raise RuntimeError('Video tải nằm ngoài thư mục dự án.')
                info=probe(video)
                if info.get('duration',0)<=0:raise RuntimeError('Video không có thời lượng hợp lệ.')
                sha=digest(video)
                atomic_json(folder/'project.json',{'version':1,'type':'downloaded_video_project','source':video.name,'source_url':job['url'],'original_local_path':job.get('local_path'),'name':job['name'],'sha256':sha,'info':info})
                pack_video_project(folder)
                job.update(status='completed',project=str((folder/(folder.name+'.frameproj')).relative_to(self.root)),video=str(video.relative_to(self.root)),sha256=sha,info=info,error='')
            except Exception as exc:job.update(status='failed',error=str(exc))
            publish(job)
        self.persist();self.save_excel()

def pack_video_project(folder):
    """A single .frameproj ZIP contains the original video and portable metadata."""
    import zipfile
    folder=Path(folder);data=json.loads((folder/'project.json').read_text(encoding='utf-8'));video=folder/data['source']
    target=folder/(folder.name+'.frameproj');temp=folder/(folder.name+'.frameproj.tmp')
    try:
        with zipfile.ZipFile(temp,'w',compression=zipfile.ZIP_STORED,allowZip64=True) as archive:
            archive.write(folder/'project.json','project.json');archive.write(video,video.name)
        os.replace(temp,target)
    finally:temp.unlink(missing_ok=True)
    return target

def open_video_project(archive_path,destination):
    import zipfile,tempfile
    archive_path=Path(archive_path);destination=Path(destination).resolve();destination.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(archive_path) as z:
        member=z.getinfo('project.json')
        if member.file_size>1024*1024:raise ValueError('Thông tin dự án quá lớn.')
        data=json.loads(z.read(member));name=data.get('source','')
        if data.get('type')!='downloaded_video_project' or not re.fullmatch(r'source\.(mp4|mkv|mov|webm|avi|m4v)',name,re.I):raise ValueError('Dự án không đúng định dạng.')
        video=z.getinfo(name)
        if video.file_size>20*1024**3:raise ValueError('Video dự án lớn hơn giới hạn 20 GB.')
        with tempfile.TemporaryDirectory(dir=destination) as temp:
            local=Path(temp)/name
            with z.open(video) as source,local.open('wb') as out:
                import shutil
                shutil.copyfileobj(source,out,1024*1024)
            if digest(local)!=data.get('sha256'):raise ValueError('Video trong dự án bị hỏng hoặc khác bản gốc.')
            stem='Imported_'+safe_name(archive_path.stem);target=destination/stem;n=1
            while target.exists():target=destination/f'{stem}_{n}';n+=1
            target.mkdir();os.replace(local,target/name);atomic_json(target/'project.json',data)
    return target/name
