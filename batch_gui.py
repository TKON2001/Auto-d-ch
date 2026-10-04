"""Excel batch download dialog; no OCR or translation dependency during downloads."""
from pathlib import Path
from PySide6.QtCore import QThread,Signal,Qt
from PySide6.QtGui import QColor
from PySide6.QtWidgets import QWidget,QCheckBox,QVBoxLayout,QHBoxLayout,QLabel,QPushButton,QLineEdit,QFileDialog,QComboBox,QTableWidget,QTableWidgetItem,QPlainTextEdit,QMessageBox,QHeaderView
from batch_core import BatchQueue
from pause_control import PauseController

class DownloadWorker(QThread):
    progress=Signal(dict);log=Signal(str);error=Signal(str);state=Signal(str)
    def __init__(self,queue):
        super().__init__();self.queue=queue;self.control=PauseController(self.state.emit)
    def run(self):
        try:self.queue.run(self.control,self.progress.emit,self.log.emit)
        except Exception as exc:self.error.emit(str(exc))

class BatchPanel(QWidget):
    video_selected=Signal(str);running_changed=Signal(bool);queue_ready=Signal(object)
    def __init__(self,parent=None):
        super().__init__(parent);self.setWindowTitle('Nhập video từ Excel • Dự án hàng loạt');self.queue=None;self.worker=None;self.setting_excel=False
        lay=QVBoxLayout(self)
        title=QLabel('NHẬP VIDEO THEO THỨ TỰ EXCEL');title.setStyleSheet('font-size:16px;font-weight:600;color:#a7e5d3');lay.addWidget(title)
        desc=QLabel('Cột Link: đường dẫn file trên máy hoặc URL. STT và Tên tùy chọn. Có STT: tăng dần. Không có STT: từ trên xuống.\nHoàn tất: xanh • Lỗi: đỏ • Mỗi video có thư mục riêng. Excel gốc giữ nguyên.');desc.setWordWrap(True);lay.addWidget(desc)
        self.excel=QLineEdit();self.excel.setPlaceholderText('File .xlsx hoặc CSV UTF-8');self.sheet=QComboBox();self.sheet.setMinimumWidth(90)
        self.dest=QLineEdit(str(Path.home()/'Videos'));self.cookies=QLineEdit();self.cookies.setPlaceholderText('Cookies Netscape .txt, tùy chọn (không lưu vào dự án)')
        self.inputs=[]
        for label,edit,handler in [('Excel',self.excel,self.pick_excel),('Lưu trên ổ',self.dest,self.pick_dest),('Cookies',self.cookies,self.pick_cookies)]:
            row=QHBoxLayout();row.addWidget(QLabel(label));row.addWidget(edit,1);b=QPushButton('Chọn…');b.clicked.connect(handler);row.addWidget(b);lay.addLayout(row);self.inputs.extend([edit,b])
        row=QHBoxLayout();row.addWidget(QLabel('Sheet'));row.addWidget(self.sheet);self.inspect=QPushButton('Nạp danh sách');self.inspect.clicked.connect(self.load_queue);row.addWidget(self.inspect);self.template=QPushButton('Lưu mẫu CSV');self.template.clicked.connect(self.save_template);lay.addLayout(row);row=QHBoxLayout();row.addWidget(self.template);self.import_project=QPushButton('Mở .frameproj');self.import_project.clicked.connect(self.open_project);row.addWidget(self.import_project);row.addStretch();lay.addLayout(row);self.inputs.extend([self.sheet,self.inspect,self.template,self.import_project])
        self.use_list=QCheckBox('Bắt đầu dịch sẽ dịch toàn bộ danh sách theo STT');lay.addWidget(self.use_list)
        self.table=QTableWidget(0,5);self.table.setHorizontalHeaderLabels(['STT / dòng','Tên video','Link','Trạng thái','Video / lỗi']);self.table.setSelectionBehavior(QTableWidget.SelectRows);self.table.setEditTriggers(QTableWidget.NoEditTriggers);self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents);self.table.horizontalHeader().setSectionResizeMode(2,QHeaderView.Stretch);self.table.setMinimumHeight(180);self.table.setMaximumHeight(230);lay.addWidget(self.table)
        row=QHBoxLayout();self.start=QPushButton('Bắt đầu nhập');self.start.setObjectName('primary');self.start.clicked.connect(self.start_download);self.pause=QPushButton('Tạm dừng');self.pause.clicked.connect(self.toggle_pause);self.stop=QPushButton('Dừng sau video này');self.stop.clicked.connect(self.stop_after);self.save=QPushButton('Lưu Excel');self.save.clicked.connect(self.save_excel);self.open=QPushButton('Xem video đã chọn');self.open.clicked.connect(self.select_video)
        for b in [self.start,self.pause,self.stop]:row.addWidget(b)
        lay.addLayout(row);row=QHBoxLayout()
        for b in [self.save,self.open]:row.addWidget(b)
        self.pause.setEnabled(False);self.stop.setEnabled(False);self.save.setEnabled(False);self.open.setEnabled(False);lay.addLayout(row)
        self.status=QLabel('Chọn Excel, chọn thư mục rồi nạp danh sách.');self.status.setWordWrap(True);lay.addWidget(self.status)
        self.sheet.currentTextChanged.connect(self.refresh_queue);self.dest.editingFinished.connect(self.refresh_queue);self.excel.editingFinished.connect(self.refresh_queue)
        self.logs=QPlainTextEdit();self.logs.setReadOnly(True);self.logs.setMaximumBlockCount(250);self.logs.setMaximumHeight(90);lay.addWidget(self.logs)
    def open_project(self):
        p,_=QFileDialog.getOpenFileName(self,'Mở dự án chứa video','','Frame Project (*.frameproj)')
        if not p:return
        destination=QFileDialog.getExistingDirectory(self,'Giải nén dự án vào thư mục',self.dest.text())
        if not destination:return
        try:
            from batch_core import open_video_project
            video=open_video_project(p,destination);self.video_selected.emit(str(video))
        except Exception as exc:QMessageBox.warning(self,'Không mở được dự án',str(exc))
    def pick_excel(self):
        p,_=QFileDialog.getOpenFileName(self,'Danh sách video','','Excel (*.xlsx);;CSV (*.csv)')
        if p:
            self.setting_excel=True
            self.excel.setText(p);self.sheet.clear()
            if Path(p).suffix.lower()=='.xlsx':
                try:
                    from openpyxl import load_workbook
                    wb=load_workbook(p,read_only=True);self.sheet.addItems(wb.sheetnames);wb.close()
                except Exception as exc:QMessageBox.warning(self,'Không đọc được Excel',str(exc))
            self.setting_excel=False;self.load_queue()
    def refresh_queue(self,*_):
        if not self.setting_excel and not (self.worker and self.worker.isRunning()) and Path(self.excel.text()).is_file():self.load_queue()
    def pick_dest(self):
        p=QFileDialog.getExistingDirectory(self,'Thư mục chứa dự án',self.dest.text())
        if p:self.dest.setText(p);self.refresh_queue()
    def pick_cookies(self):
        p,_=QFileDialog.getOpenFileName(self,'Cookies Netscape','','Text (*.txt)')
        if p:self.cookies.setText(p)
    def save_template(self):
        p,_=QFileDialog.getSaveFileName(self,'Mẫu mở bằng Excel','Danh_sach_video.csv','CSV (*.csv)')
        if p:
            try:Path(p).write_text('STT,Link,Tên\n1,D:\\Video\\video01.mp4,Video 1\n2,D:\\Video\\video02.mp4,Video 2\n',encoding='utf-8-sig');self.status.setText('Đã lưu mẫu. Điền link rồi lưu thành .xlsx hoặc CSV UTF-8.')
            except OSError as exc:QMessageBox.warning(self,'Không lưu được mẫu',str(exc))
    def load_queue(self):
        try:
            if not Path(self.excel.text()).is_file():raise ValueError('Chọn file Excel có thật.')
            if not self.dest.text().strip():raise ValueError('Chọn thư mục đích.')
            self.queue=BatchQueue(self.excel.text(),self.dest.text(),self.sheet.currentText() or None);self.table.setRowCount(len(self.queue.jobs))
            for j in self.queue.jobs:self.update_row(j)
            self.status.setText(f'{len(self.queue.jobs)} video • Dự án: {self.queue.root}');self.save.setEnabled(True);self.open.setEnabled(True);self.queue_ready.emit(self.queue)
        except Exception as exc:self.queue=None;self.use_list.setChecked(False);QMessageBox.warning(self,'Không nạp được danh sách',str(exc))
    def update_row(self,job):
        if not self.queue:return
        i=next((i for i,j in enumerate(self.queue.jobs) if j['id']==job['id']),None)
        if i is None:return
        state={'completed':'Đã tải','failed':'Lỗi','downloading':'Đang xử lý','pending':'Chờ'}.get(job['status'],job['status'])
        if job.get('local_path') and job['status']=='completed':state='Đã nhập'
        translated=job.get('translation_status')
        if translated and job['status']=='completed':state={'translating':'Đang dịch','completed':'Đã dịch','failed':'Lỗi dịch'}.get(translated,state)
        values=[f"{job['order'] if job['order'] is not None else '?'} / dòng {job['row']}",job['name'],job['url'] or job['raw'],state,job.get('translation_error') or job.get('error') or job.get('translation_output') or job.get('video','')]
        for col,value in enumerate(values):
            item=QTableWidgetItem(str(value));item.setToolTip(str(value))
            if job['status'] in ('completed','failed'):
                item.setBackground(QColor('#652c35' if job['status']=='failed' or translated=='failed' else ('#6d5b25' if translated=='translating' else '#1f553d')));item.setForeground(QColor('#ffffff'))
            self.table.setItem(i,col,item)
        complete=sum(j['status']=='completed' for j in self.queue.jobs);fail=sum(j['status']=='failed' for j in self.queue.jobs)
        self.status.setText(f'Đã tải {complete}/{len(self.queue.jobs)} • Lỗi {fail} • {state}: {job["name"]}')
    def set_running(self,running):
        self.running_changed.emit(running)
        for w in self.inputs:w.setEnabled(not running)
        self.start.setEnabled(not running);self.pause.setEnabled(running);self.stop.setEnabled(running);self.save.setEnabled(not running and self.queue is not None);self.open.setEnabled(not running and self.queue is not None)
    def start_download(self):
        if self.worker and self.worker.isRunning():return
        self.load_queue()
        if not self.queue:return
        cookies=self.cookies.text().strip()
        if cookies and not Path(cookies).is_file():QMessageBox.warning(self,'Cookies','Không tìm thấy file cookies.');return
        self.queue.cookies=cookies or None;self.logs.clear();self.set_running(True);self.pause.setText('Tạm dừng');self.worker=DownloadWorker(self.queue);self.worker.progress.connect(self.update_row);self.worker.log.connect(self.logs.appendPlainText);self.worker.error.connect(self.report_error);self.worker.state.connect(self.pause_state);self.worker.finished.connect(self.finished);self.worker.start()
    def report_error(self,error):self.logs.appendPlainText('LỖI: '+error)
    def toggle_pause(self):
        if not self.worker:return
        if self.worker.control.requested:self.worker.control.resume()
        else:self.worker.control.pause()
    def pause_state(self,state):
        self.pause.setText('Tiếp tục' if state in ('paused','pausing') else 'Tạm dừng')
        if state=='paused':self.status.setText('Đã tạm dừng. Bấm Tiếp tục để tải tiếp.')
        if state.startswith('pause_error:'):self.logs.appendPlainText(state)
    def stop_after(self):
        if self.worker:self.queue.stop_requested=True;self.worker.control.resume();self.stop.setEnabled(False);self.status.setText('Sẽ dừng sau khi video hiện tại hoàn tất.')
    def finished(self):
        self.set_running(False);self.pause.setText('Tạm dừng');complete=sum(j['status']=='completed' for j in self.queue.jobs);failed=sum(j['status']=='failed' for j in self.queue.jobs);pending=len(self.queue.jobs)-complete-failed
        self.status.setText(f'Đã tải {complete} • Lỗi {failed} • Chờ {pending}\nDự án: {self.queue.root}')
        if self.queue.warning:self.logs.appendPlainText(self.queue.warning)
        self.queue_ready.emit(self.queue)
    def save_excel(self):
        if not self.queue:return
        try:
            warning=self.queue.save_excel()
            if warning:QMessageBox.warning(self,'Chưa lưu được Excel',warning)
            else:self.status.setText('Đã lưu Excel tô màu: '+str(self.queue.result))
        except Exception as exc:QMessageBox.warning(self,'Không lưu được Excel',str(exc))
    def select_video(self):
        i=self.table.currentRow()
        if self.queue is None or i<0:return
        job=self.queue.jobs[i];p=self.queue.root/job.get('video','')
        if job['status']!='completed' or not p.is_file():QMessageBox.information(self,'Chọn video','Chọn dòng đã tải thành công.');return
        self.video_selected.emit(str(p))
    def closeEvent(self,event):
        if self.worker and self.worker.isRunning():QMessageBox.information(self,'Đang tải','Bấm Dừng sau video này và chờ tác vụ kết thúc trước khi đóng.');event.ignore()
        else:event.accept()

# Legacy import kept for existing tests; this is now an inline QWidget.
BatchDialog=BatchPanel
