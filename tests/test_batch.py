import json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from batch_core import BatchQueue,read_queue,safe_name,extract_url
from pause_control import PauseController
class Tests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.file=self.root/'links.csv';self.file.write_text('STT,Link,Tên\n2,https://example.com/b,B\n1,https://example.com/a,A\n3,invalid,Bad\n4,https://example.com/fail,Fail\n',encoding='utf-8-sig')
 def tearDown(self):self.tmp.cleanup()
 def dl(self,j,folder,*_):
  if j['name']=='Fail':raise RuntimeError('403')
  p=folder/'source.mp4';p.write_bytes(b'video'+j['name'].encode());return p
 def info(self,p):return {'duration':3,'width':320,'height':240,'has_audio':True}
 def queue(self):return BatchQueue(self.file,self.root/'output')
 def test_order(self):
  _,jobs=read_queue(self.file);self.assertEqual([j['name'] for j in jobs],['A','B','Bad','Fail']);self.assertEqual(extract_url('分享 https://www.douyin.com/?modal_id=123 。'),'https://www.douyin.com/video/123');self.assertEqual(safe_name('CON'),'_CON')
 def test_complete_color_resume_and_changed(self):
  q=self.queue();q.run(PauseController(),downloader=self.dl,probe=self.info);self.assertEqual([j['status'] for j in q.jobs],['completed','completed','failed','failed']);self.assertTrue((q.root/q.jobs[0]['folder']/'project.json').exists())
  from openpyxl import load_workbook
  wb=load_workbook(q.result);ws=wb.active;self.assertEqual(ws.cell(3,1).fill.fgColor.rgb,'00C6EFCE');self.assertEqual(ws.cell(4,1).fill.fgColor.rgb,'00FFC7CE');wb.close()
  calls=[];q=self.queue();q.run(PauseController(),downloader=lambda *a:(calls.append(a[0]['name']) or self.dl(*a)),probe=self.info);self.assertEqual(calls,['Fail'])
  (q.root/q.jobs[0]['video']).write_bytes(b'changed');calls.clear();q=self.queue();q.run(PauseController(),downloader=lambda *a:(calls.append(a[0]['name']) or self.dl(*a)),probe=self.info);self.assertEqual(calls,['A','Fail']);self.assertTrue(self.file.read_text(encoding='utf-8-sig').startswith('STT,Link,Tên'))
 def test_stop(self):
  q=self.queue()
  def dl(*a):q.stop_requested=True;return self.dl(*a)
  q.run(PauseController(),downloader=dl,probe=self.info);self.assertEqual([j['status'] for j in q.jobs],['completed','pending','pending','pending'])
 def test_excel_locked(self):
  q=self.queue();import batch_core;original=batch_core.os.replace
  def lock(a,b):
   if Path(b).suffix=='.xlsx':raise PermissionError('locked')
   return original(a,b)
  with patch('batch_core.os.replace',side_effect=lock):q.run(PauseController(),downloader=self.dl,probe=self.info)
  self.assertTrue(q.warning);self.assertEqual(json.loads(q.manifest.read_text())['videos'][0]['status'],'completed');self.assertEqual(q.save_excel(),'')
 def test_duplicate_and_invalid_stt(self):
  self.file.write_text('STT,Link,Tên\n1,https://example.com/a,../../A\n1,https://example.com/b,../../A\n,https://example.com/c,C\n',encoding='utf-8-sig');q=self.queue();self.assertEqual(len(set(j['folder'] for j in q.jobs)),3);q.run(PauseController(),downloader=self.dl,probe=self.info);self.assertEqual(q.jobs[2]['status'],'failed')
 def test_outside_rejected(self):
  p=self.root/'outside.mp4';p.write_bytes(b'x');q=self.queue();q.run(PauseController(),downloader=lambda *_:p,probe=self.info);self.assertTrue(all(j['status']=='failed' for j in q.jobs))
 def test_xlsx_and_project_roundtrip(self):
  q=self.queue();q.run(PauseController(),downloader=self.dl,probe=self.info)
  from batch_core import open_video_project,digest
  video=open_video_project(q.root/q.jobs[0]['project'],self.root/'import');self.assertEqual(digest(video),q.jobs[0]['sha256'])
  x=BatchQueue(q.result,self.root/'xlsx');x.run(PauseController(),downloader=self.dl,probe=self.info);self.assertEqual(x.jobs[0]['status'],'completed')
 def test_gui(self):
  from PySide6.QtWidgets import QApplication
  from batch_gui import BatchDialog
  app=QApplication.instance() or QApplication([]);d=BatchDialog();d.excel.setText(str(self.file));d.dest.setText(str(self.root/'out'));d.load_queue();self.assertEqual(d.table.rowCount(),4);d.queue.run(PauseController(),update=d.update_row,downloader=self.dl,probe=self.info);self.assertEqual(d.table.item(0,3).text(),'Đã tải');self.assertEqual(d.table.item(0,3).background().color().name(),'#1f553d');d.pause_state('paused');self.assertEqual(d.pause.text(),'Tiếp tục');d.close()
if __name__=='__main__':unittest.main()
