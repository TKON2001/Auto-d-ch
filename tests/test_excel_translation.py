import unittest,tempfile,json,threading,time
from pathlib import Path
from unittest.mock import patch
from batch_core import BatchQueue
from batch_translation import translate_queue
from pause_control import PauseController
class TranslationTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name);self.video=self.root/'video.mp4';self.video.write_bytes(b'video original');self.file=self.root/'list.csv';self.file.write_text(f'STT,Link,Tên\n2,{self.video},Second\n1,{self.video},First\n3,{self.video},Fail\n',encoding='utf-8-sig');self.calls=[]
 def tearDown(self):self.tmp.cleanup()
 def process(self,source_value,output,**kwargs):
  folder=Path(source_value).parent;self.calls.append(folder.name)
  if '_Fail_' in folder.name:raise RuntimeError('translation failed')
  Path(output).write_bytes(b'translated');return output
 def info(self,*_):return {'duration':2,'width':320,'height':240,'has_audio':True}
 def test_order_failure_persistence_and_resume(self):
  q=BatchQueue(self.file,self.root/'out')
  with patch('video_translator.probe',side_effect=self.info):r=translate_queue(q,{'api_key':'NOT_REAL','model':'mock'},PauseController(),self.process)
  self.assertEqual((r['success'],r['failed']),(2,1));self.assertIn('First',self.calls[0]);self.assertIn('Second',self.calls[1]);self.assertEqual(q.jobs[2]['translation_status'],'failed');self.assertNotIn('NOT_REAL',q.manifest.read_text())
  self.calls.clear();q=BatchQueue(self.file,self.root/'out')
  with patch('video_translator.probe',side_effect=self.info):r=translate_queue(q,{'api_key':'NOT_REAL','model':'mock'},PauseController(),self.process)
  self.assertEqual(len(self.calls),1);self.assertIn('Fail',self.calls[0])
  from openpyxl import load_workbook
  wb=load_workbook(q.result);self.assertEqual(wb.active.cell(4,1).fill.fgColor.rgb,'00FFC7CE');wb.close()
 def test_changed_settings_retranslate(self):
  q=BatchQueue(self.file,self.root/'out')
  with patch('video_translator.probe',side_effect=self.info):translate_queue(q,{'model':'a'},PauseController(),self.process);self.calls.clear();translate_queue(q,{'model':'b'},PauseController(),self.process)
  self.assertEqual(len(self.calls),3)
 def test_gui_load_links_to_main_start(self):
  from video_translator import launch_gui
  app,w=launch_gui(test_mode=True);w.batch.excel.setText(str(self.file));w.batch.dest.setText(str(self.root/'gui'));w.batch.load_queue();self.assertEqual(w.source.text(),str(self.video));self.assertTrue(w.batch.use_list.isChecked());self.assertIn('3 video',w.start.text());w.api.setText('NOT_REAL')
  with patch('video_translator.probe',side_effect=self.info),patch('video_translator.process',side_effect=self.process):
   w.start_job();self.assertFalse(w.start.isEnabled());limit=time.monotonic()+5
   while w.worker.isRunning() and time.monotonic()<limit:app.processEvents();time.sleep(.01)
   self.assertFalse(w.worker.isRunning());app.processEvents();self.assertTrue(w.start.isEnabled());self.assertEqual(len(self.calls),3);self.assertIn('Đã dịch 2/3',w.status.text())
  w.close()
if __name__=='__main__':unittest.main()
